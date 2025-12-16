"""
Text Extraction Service
Extracts text from office files (DOCX, XLSX, CSV, TXT) without OCR
"""
from typing import Dict, Optional
from io import BytesIO
import csv
import chardet


class TextExtractionService:
    """Service for extracting text from office and text files"""
    
    @staticmethod
    async def extract_text(mime: str, file_bytes: bytes) -> Dict[str, Optional[str]]:
        """
        Extract text from office/text files
        
        Args:
            mime: MIME type of the file
            file_bytes: File content as bytes
            
        Returns:
            {
                "text": str,
                "error": Optional[str],
                "provider": "native"
            }
        """
        try:
            if mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
                return await TextExtractionService.extract_from_docx(file_bytes)
            elif mime == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
                return await TextExtractionService.extract_from_xlsx(file_bytes)
            elif mime in ["text/csv", "application/csv"]:
                return await TextExtractionService.extract_from_csv(file_bytes)
            elif mime == "text/plain":
                return await TextExtractionService.extract_from_txt(file_bytes)
            else:
                return {
                    "text": "",
                    "error": f"Unsupported MIME type for text extraction: {mime}",
                    "provider": "native"
                }
        except Exception as e:
            return {
                "text": "",
                "error": f"Failed to extract text: {str(e)}",
                "provider": "native"
            }
    
    @staticmethod
    async def extract_from_docx(file_bytes: bytes) -> Dict[str, Optional[str]]:
        """Extract text from DOCX file"""
        try:
            from docx import Document
            
            doc = Document(BytesIO(file_bytes))
            text_parts = []
            
            # Extract text from paragraphs
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_parts.append(paragraph.text.strip())
            
            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        if cell.text.strip():
                            row_text.append(cell.text.strip())
                    if row_text:
                        text_parts.append(" | ".join(row_text))
            
            text = "\n".join(text_parts)
            
            return {
                "text": text,
                "error": None,
                "provider": "native"
            }
        except Exception as e:
            return {
                "text": "",
                "error": f"Failed to extract text from DOCX: {str(e)}",
                "provider": "native"
            }
    
    @staticmethod
    async def extract_from_xlsx(file_bytes: bytes) -> Dict[str, Optional[str]]:
        """Extract text from XLSX file"""
        try:
            from openpyxl import load_workbook
            
            workbook = load_workbook(BytesIO(file_bytes), data_only=True)
            text_parts = []
            
            # Extract text from all sheets
            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                text_parts.append(f"=== Sheet: {sheet_name} ===")
                
                # Extract text from cells
                for row in sheet.iter_rows(values_only=True):
                    row_text = []
                    for cell_value in row:
                        if cell_value is not None:
                            # Convert cell value to string
                            cell_str = str(cell_value).strip()
                            if cell_str:
                                row_text.append(cell_str)
                    if row_text:
                        text_parts.append(" | ".join(row_text))
                
                text_parts.append("")  # Empty line between sheets
            
            text = "\n".join(text_parts)
            
            return {
                "text": text,
                "error": None,
                "provider": "native"
            }
        except Exception as e:
            return {
                "text": "",
                "error": f"Failed to extract text from XLSX: {str(e)}",
                "provider": "native"
            }
    
    @staticmethod
    async def extract_from_csv(file_bytes: bytes) -> Dict[str, Optional[str]]:
        """Extract text from CSV file"""
        try:
            # Detect encoding
            detected = chardet.detect(file_bytes)
            encoding = detected.get('encoding', 'utf-8')
            
            # Try common encodings if detection fails
            encodings_to_try = [encoding, 'utf-8', 'utf-16', 'latin-1', 'cp1252']
            
            text_parts = []
            for enc in encodings_to_try:
                try:
                    # Decode bytes to string
                    csv_string = file_bytes.decode(enc)
                    
                    # Parse CSV
                    csv_reader = csv.reader(csv_string.splitlines())
                    for row in csv_reader:
                        # Filter out empty cells
                        row_text = [cell.strip() for cell in row if cell.strip()]
                        if row_text:
                            text_parts.append(" | ".join(row_text))
                    
                    break  # Successfully parsed, exit loop
                except (UnicodeDecodeError, csv.Error):
                    continue
            
            if not text_parts:
                # If all encodings failed, try to decode as plain text
                try:
                    text = file_bytes.decode('utf-8', errors='replace')
                    text_parts = [line.strip() for line in text.splitlines() if line.strip()]
                except:
                    pass
            
            text = "\n".join(text_parts)
            
            return {
                "text": text,
                "error": None,
                "provider": "native"
            }
        except Exception as e:
            return {
                "text": "",
                "error": f"Failed to extract text from CSV: {str(e)}",
                "provider": "native"
            }
    
    @staticmethod
    async def extract_from_txt(file_bytes: bytes) -> Dict[str, Optional[str]]:
        """Extract text from TXT file"""
        try:
            # Detect encoding
            detected = chardet.detect(file_bytes)
            encoding = detected.get('encoding', 'utf-8')
            
            # Try common encodings if detection fails
            encodings_to_try = [encoding, 'utf-8', 'utf-16', 'latin-1', 'cp1252']
            
            text = ""
            for enc in encodings_to_try:
                try:
                    text = file_bytes.decode(enc)
                    break  # Successfully decoded, exit loop
                except UnicodeDecodeError:
                    continue
            
            if not text:
                # If all encodings failed, use replace mode
                text = file_bytes.decode('utf-8', errors='replace')
            
            return {
                "text": text,
                "error": None,
                "provider": "native"
            }
        except Exception as e:
            return {
                "text": "",
                "error": f"Failed to extract text from TXT: {str(e)}",
                "provider": "native"
            }

